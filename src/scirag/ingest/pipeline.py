from __future__ import annotations
import json
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional

from ..db import get_conn
from ..embed import embed_texts, to_blob

TIER_CHUNK_SIZES = {1: 200, 2: 500, 3: 300, 4: 800}
TIER_MODES = {1: "divulgacao", 2: "investigacao", 3: "pedagogico", 4: "investigacao"}
BATCH = 64

# ── Standalone chunker (no fdllmret / no OpenAI dependency) ──────────────────
import tiktoken as _tiktoken

_tokenizer = _tiktoken.get_encoding("cl100k_base")
_MIN_CHARS = 350
_MIN_LEN = 5
_MAX_CHUNKS = 10000


def _extract_text(path: Path) -> Optional[str]:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        try:
            import fitz
            doc = fitz.open(str(path))
            return "\n".join(page.get_text() for page in doc)
        except Exception:
            from pypdf import PdfReader
            return "\n".join(p.extract_text() or "" for p in PdfReader(str(path)).pages)
    elif suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            parts = [f"[Slide {i}]"]
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
            slides.append("\n".join(parts))
        return "\n\n".join(slides)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(str(path))
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return None


def _chunk_text(text: str, chunk_size: int) -> List[str]:
    if not text or text.isspace():
        return []
    tokens = _tokenizer.encode(text, disallowed_special=())
    chunks, n = [], 0
    while tokens and n < _MAX_CHUNKS:
        chunk = tokens[:chunk_size]
        chunk_text = _tokenizer.decode(chunk)
        if not chunk_text or chunk_text.isspace():
            tokens = tokens[len(chunk):]
            continue
        cut = max(chunk_text.rfind("."), chunk_text.rfind("?"),
                  chunk_text.rfind("!"), chunk_text.rfind("\n"))
        if cut != -1 and cut > _MIN_CHARS:
            chunk_text = chunk_text[:cut + 1]
        ct = chunk_text.replace("\n", " ").strip()
        if len(ct) > _MIN_LEN:
            chunks.append(ct)
        tokens = tokens[len(_tokenizer.encode(chunk_text, disallowed_special=())):]
        n += 1
    if tokens:
        tail = _tokenizer.decode(tokens).replace("\n", " ").strip()
        if len(tail) > _MIN_LEN:
            chunks.append(tail)
    return chunks


def ingest_folder(
    folder: str | Path,
    tier: int,
    db_path: Optional[Path] = None,
    exts: tuple = (".pdf", ".txt", ".md", ".pptx", ".docx"),
    skip_existing: bool = True,
) -> int:
    folder = Path(folder)
    chunk_size = TIER_CHUNK_SIZES.get(tier, 300)
    writing_mode = TIER_MODES.get(tier, "investigacao")
    count = 0

    with get_conn(db_path) as conn:
        for fpath in sorted(folder.rglob("*")):
            if fpath.name.startswith("~$"):  # skip Word temp lock files
                continue
            if fpath.suffix.lower() not in exts:
                continue

            source_id = str(uuid.uuid5(uuid.NAMESPACE_URL, fpath.as_posix()))

            if skip_existing and conn.execute(
                "SELECT 1 FROM sources WHERE id = ?", (source_id,)
            ).fetchone():
                continue

            if fpath.stat().st_size == 0:
                print(f"  skipping empty file: {fpath.name}")
                continue

            try:
                text = _extract_text(fpath)
            except Exception as exc:
                print(f"  skipping unreadable file: {fpath.name} ({exc})")
                continue
            if not text or text.isspace():
                continue

            now = datetime.now(timezone.utc).isoformat()
            conn.execute(
                """INSERT OR REPLACE INTO sources
                   (id, title, source_type, language, tier, ingested_at, full_text, metadata)
                   VALUES (?, ?, 'file', 'pt', ?, ?, ?, ?)""",
                (source_id, fpath.stem, tier, now, text,
                 json.dumps({"file_path": str(fpath)})),
            )

            chunks = _chunk_text(text, chunk_size)
            if not chunks:
                conn.commit()
                continue

            node_rows: list[tuple] = []
            for i in range(0, len(chunks), BATCH):
                batch_texts = chunks[i: i + BATCH]
                vecs = embed_texts(batch_texts)
                for j, (ct, vec) in enumerate(zip(batch_texts, vecs)):
                    node_rows.append((
                        f"{source_id}_{i + j}",
                        ct, fpath.stem, writing_mode,
                        source_id, tier, now, to_blob(vec),
                    ))

            conn.executemany(
                """INSERT OR REPLACE INTO nodes
                   (id, content, title, writing_mode, source_id, tier, created_at, embedding)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?)""",
                node_rows,
            )
            conn.commit()
            count += 1
            print(f"  ingested {fpath.name} → {len(node_rows)} nodes")

    return count
